from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from chapter1_docx_generator import (
    add_body_paragraph,
    add_caption_label,
    add_caption_title,
    add_centered_paragraph,
    add_section_heading,
    apply_font_to_paragraph,
    apply_table_grid,
    configure_section,
    configure_section_headers,
    format_paragraph,
    set_run_font,
    shade_cell,
)


DOCS_DIR = Path(__file__).parent
DEFAULT_BASE_PATH = DOCS_DIR / "MARTINEZ_LLAWDERES_ARGA_CHAPTERS_1_AND_2_REVISED.docx"
DEFAULT_SOURCE_PATH = DOCS_DIR / "Chapter 3 MDs" / "chapter3_tables_draft.md"
DEFAULT_OUTPUT_PATH = DOCS_DIR / "MARTINEZ_LLAWDERES_ARGA_CHAPTERS_1_TO_3_TABLES_DRAFT.docx"


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def parse_markdown_row(line: str) -> list[str]:
    content = line.strip().strip("|")
    cells = [cell.strip() for cell in content.split("|")]
    return [cell.replace("`", "") for cell in cells]


def iter_markdown_tokens(source_path: Path):
    lines = source_path.read_text(encoding="utf-8").splitlines()
    i = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer).strip()
            paragraph_buffer = []
            if text:
                return ("p", text)
        return None

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if not line:
            token = flush_paragraph()
            if token:
                yield token
            i += 1
            continue

        if line in {"# CHAPTER 3", "# METHODOLOGY"}:
            token = flush_paragraph()
            if token:
                yield token
            i += 1
            continue

        if line.startswith("## "):
            token = flush_paragraph()
            if token:
                yield token
            yield ("h2", line[3:].strip())
            i += 1
            continue

        if line.startswith("### "):
            token = flush_paragraph()
            if token:
                yield token
            yield ("h3", line[4:].strip())
            i += 1
            continue

        match = re.fullmatch(r"(Table|Figure)\s+(\d+)", line)
        if match:
            token = flush_paragraph()
            if token:
                yield token
            caption_type = match.group(1)
            number = int(match.group(2))
            title = ""
            j = i + 1
            while j < len(lines):
                candidate = lines[j].strip()
                if candidate:
                    title = candidate
                    break
                j += 1
            yield ("caption", (caption_type, number, title))
            i = j + 1
            continue

        image_match = re.fullmatch(r"!\[[^\]]*\]\(([^)]+)\)", line)
        if image_match:
            token = flush_paragraph()
            if token:
                yield token
            yield ("image", image_match.group(1).strip())
            i += 1
            continue

        if line.startswith("|"):
            token = flush_paragraph()
            if token:
                yield token
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = parse_markdown_row(lines[i])
                if row and not is_separator_row(row):
                    rows.append(row)
                i += 1
            yield ("table", rows)
            continue

        paragraph_buffer.append(line)
        i += 1

    token = flush_paragraph()
    if token:
        yield token


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_fixed_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_table_grid_widths(table, widths_inches: list[float]) -> None:
    total_width = sum(widths_inches)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(total_width * 1440)))
    tbl_w.set(qn("w:type"), "dxa")

    existing_grid = table._tbl.find(qn("w:tblGrid"))
    if existing_grid is not None:
        table._tbl.remove(existing_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_inches:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    table._tbl.insert(1, grid)


def repeat_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_cell_margins(table, margin_twips: int = 70) -> None:
    tbl_pr = table._tbl.tblPr
    cell_margin = tbl_pr.find(qn("w:tblCellMar"))
    if cell_margin is None:
        cell_margin = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_margin)

    for edge in ("top", "left", "bottom", "right"):
        node = cell_margin.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            cell_margin.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized_rows), cols=column_count)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_cell_margins(table)
    apply_table_grid(table)

    if column_count == 4:
        widths = [1.45, 1.15, 1.55, 2.05]
        font_size = 7.2
    elif column_count == 3:
        widths = [1.55, 1.55, 3.10]
        font_size = 9.0
    else:
        widths = [6.2 / column_count] * column_count
        font_size = 8.5

    for row_index, row in enumerate(normalized_rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = format_table_cell_text(value, col_index, column_count, row_index == 0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

            if row_index == 0:
                shade_cell(cell, "D9E2F3")

            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = None
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = 0
                paragraph.paragraph_format.space_after = 0
                if paragraph.runs:
                    paragraph.runs[0].bold = row_index == 0
                apply_font_to_paragraph(paragraph, size=font_size)

    repeat_header_row(table.rows[0])

    spacer = doc.add_paragraph()
    format_paragraph(
        spacer,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_after=3,
    )


def add_data_dictionary_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return

    for row in rows[1:]:
        normalized = row + [""] * (4 - len(row))
        field, dtype, constraints, description = normalized[:4]
        paragraph = doc.add_paragraph()
        field_run = paragraph.add_run(field)
        set_run_font(field_run, size=9, bold=True)
        detail_run = paragraph.add_run(
            f" - Type: {dtype}; Constraints: {constraints}; Description: {description}"
        )
        set_run_font(detail_run, size=9)
        format_paragraph(
            paragraph,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_indent=None,
            line_spacing=1.0,
            space_before=0,
            space_after=2,
        )

    spacer = doc.add_paragraph()
    format_paragraph(
        spacer,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_after=3,
    )


def format_table_cell_text(value: str, col_index: int, column_count: int, is_header: bool) -> str:
    if is_header or column_count != 4:
        return value

    text = value
    if col_index == 0:
        text = text.replace("_", "_\n")
        if len(text) > 18 and "\n" not in text:
            text = "\n".join(textwrap_chunks(text, 12))
    elif col_index == 1:
        replacements = {
            "timestamptz": "timestamp",
            "decimal(10,2)": "decimal\n(10,2)",
            "numeric(10,7)": "numeric\n(10,7)",
            "gen_random_uuid()": "gen_random_uuid\n()",
        }
        for old, new in replacements.items():
            text = text.replace(old, new)
    elif col_index == 2:
        text = text.replace("Primary Key", "PK")
        text = text.replace("Foreign Key", "FK")
        text = text.replace("Not Null", "NOT NULL")
        text = text.replace(", ", ",\n")
        text = text.replace(" to ", " to\n")
        text = text.replace(" Default ", "\nDefault ")
    return text


def textwrap_chunks(text: str, width: int) -> list[str]:
    return [text[index : index + width] for index in range(0, len(text), width)]


def add_markdown_image(doc: Document, source_path: Path, image_reference: str) -> None:
    image_path = (source_path.parent / image_reference).resolve()
    width_inches = 6.2
    try:
        from PIL import Image

        with Image.open(image_path) as image:
            image_width, image_height = image.size
            max_width = 6.2
            max_height = 4.9 if "screenshots" in str(image_path).lower() else 6.35
            width_scale = max_width / image_width
            height_scale = max_height / image_height
            scale = min(width_scale, height_scale)
            width_inches = image_width * scale
    except Exception:
        width_inches = 6.2

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=None,
        line_spacing=1.0,
        space_after=6,
    )


def add_minor_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=11, bold=True)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=6,
        space_after=3,
    )


def append_chapter_three(doc: Document, source_path: Path) -> None:
    chapter_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(chapter_section)
    configure_section_headers(chapter_section, show_page_number=True, hide_first_page=True)

    add_centered_paragraph(doc, "CHAPTER 3", size=12, bold=True, after=0)
    add_centered_paragraph(doc, "METHODOLOGY", size=12, bold=True, after=12)

    in_data_dictionary = False
    for token_type, value in iter_markdown_tokens(source_path):
        if token_type == "h2":
            add_section_heading(doc, value)
        elif token_type == "h3":
            add_minor_heading(doc, value)
        elif token_type == "caption":
            caption_type, number, title = value
            if caption_type == "Table":
                if number == 2:
                    in_data_dictionary = True
                elif number == 3:
                    in_data_dictionary = False
            if caption_type == "Figure" and (6 <= number <= 8 or 54 <= number <= 56):
                doc.add_page_break()
            add_caption_label(doc, caption_type, number)
            add_caption_title(doc, title)
            if len(doc.paragraphs) >= 2:
                doc.paragraphs[-2].paragraph_format.keep_with_next = True
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
        elif token_type == "table":
            if in_data_dictionary:
                add_data_dictionary_table(doc, value)
            else:
                add_markdown_table(doc, value)
        elif token_type == "image":
            add_markdown_image(doc, source_path, value)
        else:
            add_body_paragraph(doc, value)


def build_document(base_path: Path, output_path: Path, source_path: Path) -> None:
    doc = Document(str(base_path))
    append_chapter_three(doc, source_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def main() -> None:
    parser = argparse.ArgumentParser(description="Append Chapter 3 table draft to the Chapter 1-2 manuscript.")
    parser.add_argument("--base", type=Path, default=DEFAULT_BASE_PATH)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    args = parser.parse_args()

    build_document(args.base, args.output, args.source)
    print(f"WROTE={args.output}")


if __name__ == "__main__":
    main()
