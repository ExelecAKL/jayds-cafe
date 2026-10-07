from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TEMPLATE_PATH = Path(
    r"C:\Users\Exelec\Downloads\Documents & Data\Documents\Copy of KLD-Research-Format-v1.5 FOR BSIS.docx"
)
SOURCE_PATH = Path(__file__).with_name("chapter2_revised_source.md")
DEFAULT_OUTPUT_PATH = Path(
    r"C:\Users\Exelec\Downloads\School\Capstone - Chapter 2\Chapter_2_Review_of_Related_Literature_and_Studies_REVISED.docx"
)

TITLE_LINES = [
    "THE DEVELOPMENT OF A PROGRESSIVE WEB APPLICATION-BASED",
    "ORDERING SYSTEM FOR JAYD'S CAFE",
    "(CHAPTER TWO DRAFT)",
]


def set_run_font(run, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Open Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Open Sans")
    r_fonts.set(qn("w:hAnsi"), "Open Sans")
    r_fonts.set(qn("w:eastAsia"), "Open Sans")


def format_paragraph(
    paragraph,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_indent: float | None = 0.5,
    left_indent: float | None = None,
    line_spacing: float = 2.0,
    space_before: float = 0,
    space_after: float = 0,
) -> None:
    paragraph.alignment = align
    paragraph_format = paragraph.paragraph_format
    paragraph_format.line_spacing = line_spacing
    paragraph_format.space_before = Pt(space_before)
    paragraph_format.space_after = Pt(space_after)
    paragraph_format.first_line_indent = Inches(first_line_indent) if first_line_indent is not None else None
    paragraph_format.left_indent = Inches(left_indent) if left_indent is not None else None


def clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def configure_section(section) -> None:
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        configure_section(section)

    normal = next(style for style in doc.styles if (style.name or "").lower() == "normal")
    normal.font.name = "Open Sans"
    normal.font.size = Pt(11)
    normal_r_pr = normal.element.get_or_add_rPr()
    normal_r_fonts = normal_r_pr.get_or_add_rFonts()
    normal_r_fonts.set(qn("w:ascii"), "Open Sans")
    normal_r_fonts.set(qn("w:hAnsi"), "Open Sans")
    normal_r_fonts.set(qn("w:eastAsia"), "Open Sans")


def add_centered_paragraph(doc: Document, text: str, *, size: float = 12, bold: bool = False, after: float = 0) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=None,
        line_spacing=1.0,
        space_after=after,
    )


def add_heading(doc: Document, text: str, *, size: float = 12, italic: bool = False, after: float = 6) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True, italic=italic)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=6,
        space_after=after,
    )


def add_caption_label(doc: Document, caption_type: str, number: int) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"{caption_type} {number}")
    set_run_font(run, size=11, bold=True)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=6,
        space_after=0,
    )


def add_caption_title(doc: Document, title: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(title)
    set_run_font(run, size=11, italic=True)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=0,
        space_after=6,
    )


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    format_paragraph(paragraph)


def add_reference_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    url_pattern = re.compile(r"https?://\S+")
    last_index = 0

    for match in url_pattern.finditer(text):
        if match.start() > last_index:
            run = paragraph.add_run(text[last_index:match.start()])
            set_run_font(run, size=11, italic=True)

        url_run = paragraph.add_run(match.group(0))
        set_run_font(url_run, size=11, italic=True)
        url_run.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
        url_run.font.underline = True
        last_index = match.end()

    if last_index < len(text):
        run = paragraph.add_run(text[last_index:])
        set_run_font(run, size=11, italic=True)

    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=-0.5,
        left_indent=0.5,
        line_spacing=2.0,
    )


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def apply_table_grid(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_fixed_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_margins(table, margin_twips: int = 55) -> None:
    tbl_pr = table._tbl.tblPr
    cell_margin = tbl_pr.find(qn("w:tblCellMar"))
    if cell_margin is None:
        cell_margin = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_margin)

    for edge in ("top", "left", "bottom", "right"):
        node = cell_margin.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            cell_margin.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def parse_markdown_row(line: str) -> list[str]:
    content = line.strip().strip("|")
    return [cell.strip().replace("`", "") for cell in content.split("|")]


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized_rows), cols=column_count)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_table_layout(table)
    set_cell_margins(table)
    apply_table_grid(table)

    if column_count == 13:
        widths = [2.7] + [0.29] * 12
        font_size = 7.2
    else:
        widths = [6.2 / column_count] * column_count
        font_size = 8.5

    for row_index, row in enumerate(normalized_rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[col_index])

            if row_index == 0:
                shade_cell(cell, "D9E2F3")

            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if col_index == 0 and row_index > 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.first_line_indent = None
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=row_index == 0)

    spacer = doc.add_paragraph()
    format_paragraph(
        spacer,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_after=3,
    )


def add_title_page(doc: Document) -> None:
    doc.add_paragraph()
    doc.add_paragraph()
    for line in TITLE_LINES:
        add_centered_paragraph(doc, line, size=12, bold=True, after=2)

    doc.add_paragraph()
    add_centered_paragraph(doc, "AN UNDERGRADUATE CAPSTONE PROJECT", size=12, bold=True, after=12)
    add_centered_paragraph(doc, "Presented to the Faculty of the", size=11, after=0)
    add_centered_paragraph(doc, "Institute of Computing and Digital Innovations", size=11, after=0)
    add_centered_paragraph(doc, "Kolehiyo ng Lungsod ng Dasmarinas", size=11, after=0)
    add_centered_paragraph(doc, "City of Dasmarinas, Cavite, Philippines", size=11, after=12)
    add_centered_paragraph(doc, "In Partial Fulfilment", size=11, after=0)
    add_centered_paragraph(doc, "of the Requirements for the Degree", size=11, after=0)
    add_centered_paragraph(doc, "Bachelor of Science in Information Systems", size=11, after=18)
    add_centered_paragraph(doc, "AARON C. MARTINEZ", size=11, after=0)
    add_centered_paragraph(doc, "CHARMAINE G. LLAWDERES", size=11, after=0)
    add_centered_paragraph(doc, "JAMES WILFRED B. ARGA", size=11, after=18)
    add_centered_paragraph(doc, "APRIL 2026", size=11, after=0)


def clear_story_container(container) -> None:
    element = container._element
    for child in list(element):
        element.remove(child)
    container.add_paragraph()


def add_page_number_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, size=11)

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")

    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "

    field_separator = OxmlElement("w:fldChar")
    field_separator.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "1"

    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")

    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_separator)
    run._r.append(placeholder)
    run._r.append(field_end)


def set_section_page_start(section, start: int) -> None:
    section_properties = section._sectPr
    page_number_type = section_properties.find(qn("w:pgNumType"))
    if page_number_type is None:
        page_number_type = OxmlElement("w:pgNumType")
        section_properties.append(page_number_type)
    page_number_type.set(qn("w:start"), str(start))


def configure_section_headers(section, *, show_page_number: bool, hide_first_page: bool = False) -> None:
    section.different_first_page_header_footer = hide_first_page

    for story in [section.header, section.footer, section.first_page_header, section.first_page_footer]:
        story.is_linked_to_previous = False
        clear_story_container(story)

    if show_page_number:
        header_paragraph = section.header.paragraphs[0]
        format_paragraph(
            header_paragraph,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            first_line_indent=None,
            line_spacing=1.0,
        )
        add_page_number_field(header_paragraph)

    blank_header = section.first_page_header.paragraphs[0]
    format_paragraph(
        blank_header,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_line_indent=None,
        line_spacing=1.0,
    )

    for story in [section.footer, section.first_page_footer]:
        footer_paragraph = story.paragraphs[0]
        format_paragraph(
            footer_paragraph,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            first_line_indent=None,
            line_spacing=1.0,
        )


def parse_source(source_path: Path) -> list[tuple[str, str]]:
    text = source_path.read_text(encoding="utf-8")
    tokens: list[tuple[str, object]] = []
    lines = text.splitlines()
    buffer: list[str] = []
    in_references = False
    i = 0

    def flush_buffer() -> None:
        nonlocal buffer
        if not buffer:
            return
        paragraph = " ".join(part.strip() for part in buffer).strip()
        if paragraph:
            tokens.append(("ref" if in_references else "p", paragraph))
        buffer = []

    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            flush_buffer()
            i += 1
            continue

        if stripped.startswith("## "):
            flush_buffer()
            heading = stripped[3:].strip()
            tokens.append(("h2", heading))
            in_references = heading.lower() == "references"
            i += 1
            continue

        if stripped.startswith("### "):
            flush_buffer()
            tokens.append(("h3", stripped[4:].strip()))
            i += 1
            continue

        if stripped.startswith("#### "):
            flush_buffer()
            tokens.append(("h4", stripped[5:].strip()))
            i += 1
            continue

        caption_match = re.fullmatch(r"(Table|Figure)\s+(\d+)", stripped)
        if caption_match:
            flush_buffer()
            caption_type = caption_match.group(1)
            number = int(caption_match.group(2))
            title = ""
            j = i + 1
            while j < len(lines):
                candidate = lines[j].strip()
                if candidate:
                    title = candidate
                    break
                j += 1
            tokens.append(("caption", (caption_type, number, title)))
            i = j + 1
            continue

        if stripped.startswith("|"):
            flush_buffer()
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = parse_markdown_row(lines[i])
                if row and not is_separator_row(row):
                    rows.append(row)
                i += 1
            tokens.append(("table", rows))
            continue

        buffer.append(stripped)
        i += 1

    flush_buffer()
    return tokens


def body_word_count(source_path: Path) -> int:
    text = source_path.read_text(encoding="utf-8")
    body = text.split("\n## References\n", 1)[0]
    return len(re.findall(r"\b[\w'-]+\b", body))


def reference_count(source_path: Path) -> int:
    text = source_path.read_text(encoding="utf-8")
    references = text.split("\n## References\n", 1)[1]
    return len([line for line in references.splitlines() if line.strip()])


def build_document(output_path: Path, template_path: Path, source_path: Path) -> tuple[int, int]:
    doc = Document(str(template_path))
    clear_document(doc)
    configure_document(doc)

    title_section = doc.sections[0]
    configure_section_headers(title_section, show_page_number=False)
    add_title_page(doc)

    chapter_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(chapter_section)
    configure_section_headers(chapter_section, show_page_number=True, hide_first_page=True)
    set_section_page_start(chapter_section, 1)

    add_centered_paragraph(doc, "CHAPTER 2", size=12, bold=True, after=0)
    add_centered_paragraph(doc, "REVIEW OF RELATED LITERATURE AND STUDIES", size=12, bold=True, after=12)

    for token_type, value in parse_source(source_path):
        if token_type == "h2":
            add_heading(doc, value, size=12, after=6)
        elif token_type == "h3":
            add_heading(doc, value, size=11, after=4)
        elif token_type == "h4":
            add_heading(doc, value, size=11, italic=True, after=2)
        elif token_type == "caption":
            caption_type, number, title = value
            add_caption_label(doc, caption_type, number)
            add_caption_title(doc, title)
        elif token_type == "table":
            add_markdown_table(doc, value)
        elif token_type == "ref":
            add_reference_paragraph(doc, value)
        else:
            add_body_paragraph(doc, value)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return body_word_count(source_path), reference_count(source_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the revised Chapter Two .docx file.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    parser.add_argument("--template", type=Path, default=TEMPLATE_PATH)
    parser.add_argument("--source", type=Path, default=SOURCE_PATH)
    parser.add_argument("--dry-run", action="store_true", help="Print body and reference counts without writing a file.")
    args = parser.parse_args()

    if args.dry_run:
        print(f"BODY_WORD_COUNT={body_word_count(args.source)}")
        print(f"REFERENCE_COUNT={reference_count(args.source)}")
        return

    words, refs = build_document(args.output, args.template, args.source)
    print(f"BODY_WORD_COUNT={words}")
    print(f"REFERENCE_COUNT={refs}")
    print(f"WROTE={args.output}")


if __name__ == "__main__":
    main()
