from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from chapter1_docx_generator import clear_document, configure_document, format_paragraph, set_run_font


DOCS_DIR = Path(__file__).parent
TEMPLATE_CANDIDATES = [
    Path(r"C:\Users\Exelec\Downloads\Copy of KLD-Research-Format-v1.5 FOR BSIS.docx"),
    Path(r"C:\Users\Exelec\Downloads\Documents & Data\Documents\Copy of KLD-Research-Format-v1.5 FOR BSIS.docx"),
]
OUTPUT_PATH = DOCS_DIR / "JAYDS_CAFE_DFD_LEVEL_0_AND_1.docx"
LEVEL0_PATH = DOCS_DIR / "dfd_level_0_classic.png"
LEVEL1_PATH = DOCS_DIR / "dfd_level_1_classic.png"


def resolve_template() -> Path:
    for candidate in TEMPLATE_CANDIDATES:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("No KLD template copy was found.")


def configure_landscape(doc: Document) -> None:
    configure_document(doc)
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=12, bold=True)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=0,
        space_after=12,
    )


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=0.5,
        line_spacing=2.0,
        space_before=0,
        space_after=0,
    )


def add_figure_caption(doc: Document, label: str, title: str) -> None:
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=11, bold=True)
    title_run = paragraph.add_run(f" {title}")
    set_run_font(title_run, size=11, italic=True)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=8,
        space_after=4,
    )


def add_figure(doc: Document, image_path: Path) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(8.75))
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=0,
        space_after=8,
    )


def build_document() -> Path:
    template_path = resolve_template()
    doc = Document(str(template_path))
    clear_document(doc)
    configure_landscape(doc)

    add_heading(doc, "DATA FLOW DIAGRAMS OF THE PROPOSED JAYD'S CAFE ORDERING AND OPERATIONS SYSTEM")

    add_body(
        doc,
        "The Data Flow Diagram Level 0 presents the proposed Jayd's Cafe Ordering and Operations System "
        "as a single major process. It shows the main data exchanges between the system and its external "
        "entities, namely the customer, POS staff, admin or manager, Maya Payment Gateway, and Lalamove "
        "Delivery Service.",
    )
    add_figure_caption(
        doc,
        "Figure 1",
        "Data Flow Diagram Level 0 of the Proposed Jayd's Cafe Ordering and Operations System",
    )
    add_figure(doc, LEVEL0_PATH)

    doc.add_page_break()

    add_body(
        doc,
        "The Data Flow Diagram Level 1 decomposes the proposed system into its major internal processes. "
        "It illustrates how customer order requests, payment records, fulfillment updates, delivery data, "
        "menu records, loyalty information, and administrative reporting move across the system's processes "
        "and data stores.",
    )
    add_figure_caption(
        doc,
        "Figure 2",
        "Data Flow Diagram Level 1 of the Proposed Jayd's Cafe Ordering and Operations System",
    )
    add_figure(doc, LEVEL1_PATH)

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> None:
    output_path = build_document()
    print(f"WROTE={output_path}")


if __name__ == "__main__":
    main()
