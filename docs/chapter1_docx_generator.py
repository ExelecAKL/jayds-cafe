from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


TEMPLATE_PATH = Path(
    r"C:\Users\Exelec\Downloads\Documents & Data\Documents\Copy of KLD-Research-Format-v1.5 FOR BSIS.docx"
)
DEFAULT_OUTPUT_PATH = Path(
    r"C:\Users\Exelec\Downloads\Documents & Data\Documents\MARTINEZ_LLAWDERES_ARGA_C-1DRAFT_REVISED.docx"
)
FIGURE_PATH = Path(__file__).with_name("research_paradigm_jayds.png")

TITLE_LINES = [
    "THE DEVELOPMENT OF A PROGRESSIVE WEB APPLICATION-BASED",
    "ORDERING SYSTEM FOR JAYD'S CAFE",
    "(CHAPTER ONE DRAFT)",
]

BODY_SECTIONS = [
    (
        "Project Context",
        [
            "In the contemporary food and beverage industry, digital service channels have become increasingly important in shaping customer expectations. Consumers now prefer food and beverage businesses that offer convenient, fast, and reliable ordering experiences that move smoothly from mobile browsing to actual product pickup or fulfillment. For independent cafes, this shift is significant because customers no longer evaluate a business only by product quality, but also by how efficiently orders are received, prepared, and communicated. As a result, the adoption of web-based and mobile-friendly ordering solutions has become a practical strategy for improving both service delivery and customer retention.",
            "Situated within this environment is Jayd's Cafe, an independent coffee shop located in Diamond Village, Salawag, Dasmarinas, Cavite. The cafe primarily serves walk-in customers from the surrounding community while also accommodating increasing demand for advance and online-assisted ordering. At present, however, much of its daily transaction flow remains dependent on traditional counter-based service. Orders are still heavily reliant on manual communication between customer and staff, while loyalty incentives are tied to physical stamp cards. Although this setup may be workable during light traffic periods, it becomes less effective once customer volume begins to rise.",
            "The operational strain becomes more visible during afternoon peak periods, particularly from around 2:00 PM to 4:00 PM, and becomes even more pronounced during the summer season when customer demand and online-assisted orders increase. Under normal busy conditions, customers may already wait around 10 to 15 minutes for preparation, while waiting time may extend to as much as 10 to 30 minutes when multiple orders arrive at the same time. Because the order-taking process still depends on face-to-face coordination at the counter, congestion can build near the entrance while customers are lining up, placing orders, waiting for clarifications, and checking payment details.",
            "Jayd's Cafe also faces several specific bottlenecks that directly affect order accuracy and workflow efficiency. Disorganization in point-of-sale handling can slow down order capture, especially when staff must manage multiple customers simultaneously. Dependence on physical cash also creates delays whenever exact change is unavailable. In addition, uncertainty over ingredient availability can force last-minute order changes, creating confusion for both staff and customers. Manual order errors are also more likely to occur when drinks have similar names, when the temperature of the drink is not clarified, or when customers request customizations such as less sugar, hot or iced preparation, and other special remarks.",
            "Customer retention presents another important concern. While customers are generally receptive to loyalty rewards, the current physical stamp-card approach is vulnerable to loss, inconsistent stamping, and possible duplication or fraud. A paper-based system also provides little visibility to management regarding actual customer participation and repeat-purchase behavior. In this context, a digital loyalty and ordering environment becomes more than a convenience feature; it becomes a way to formalize transactions, improve customer trust, and support data-informed business decisions.",
            "Given these conditions, there is a clear rationale for developing an ordering-centered Progressive Web App system for Jayd's Cafe. A browser-based customer application can reduce entrance congestion by shifting part of the ordering process away from the counter, while barista and management support features can improve queue monitoring, stock visibility, and service coordination. By combining digital ordering, customization, third-party digital payment support, loyalty management, and administrative oversight in one coherent system, the proposed study aims to respond directly to Jayd's Cafe's operational context rather than offering a generic software solution.",
        ],
    ),
    (
        "Objectives of the Study",
        {
            "General Objective": [
                "The general objective of this study is to design and develop an ordering-centered Progressive Web App system for Jayd's Cafe that streamlines order processing, reduces queue congestion, improves order accuracy, and strengthens customer retention."
            ],
            "Specific Objectives": [
                "Evaluate the existing manual ordering process, counter workflow, loyalty practice, and stock-related bottlenecks at Jayd's Cafe.",
                "Determine the functional requirements of customers, baristas, and management based on the cafe's day-to-day operational needs.",
                "Design and develop a customer-facing Progressive Web App that supports menu browsing, beverage customization, order-ahead convenience, order-status monitoring, and third-party digital payment support through Maya Business.",
                "Build an integrated digital loyalty and promotional coupon mechanism to replace the cafe's physical stamp-card practice.",
                "Develop barista and administrative support features for order-queue visibility, menu and price management, stock awareness, and business analytics.",
                "Evaluate the developed system using ISO 25010 quality criteria and user feedback in terms of functionality, reliability, usability, security, and operational usefulness.",
            ],
        },
    ),
    (
        "Purpose and Description",
        {
            "Purpose of the Study": [
                "This study is undertaken to develop a practical digital ordering environment that responds directly to the operating conditions of Jayd's Cafe. The project is necessary because the cafe still relies heavily on counter-based order capture, physical loyalty cards, and manual coordination during peak periods, conditions that contribute to queue buildup, 10-to-30-minute waiting windows, order clarifications, and recurring customization errors. By developing an ordering-centered Progressive Web App with customer, POS, and admin support, the study aims to produce a usable intervention for the cafe while also contributing a concrete local example of applied systems development."
            ],
            "Target Company, Organization, and Institution": [
                "For Jayd's Cafe, the project is intended to reduce delays in order handling, improve the accuracy of customized drink requests, and replace the limitations of paper-based loyalty tracking. The system provides a structured digital channel for order capture, order monitoring, promotions, and selected business reporting so that the cafe can move toward a more organized and data-informed daily operation."
            ],
            "Target Office": [
                "For management and baristas, the system is designed to support clearer coordination throughout the order lifecycle. It helps staff monitor incoming orders, review customer customizations, check item availability, and follow queue movement more consistently. On the administrative side, it supports menu and price updates, promotional monitoring, and access to operational indicators that can guide weekly review and decision-making."
            ],
            "Target Client": [
                "For customers, the project offers a more convenient and transparent ordering experience. Instead of depending entirely on a physical line at the counter, customers can browse the menu, customize beverages, submit orders, track status, and access loyalty incentives through a browser-based interface. This is expected to reduce uncertainty during busy hours and improve trust in the ordering process."
            ],
            "Researchers": [
                "For the researchers, the study serves as an opportunity to apply analysis, design, development, and evaluation methods to a real business problem. It demonstrates how software engineering principles can be used to build a role-based ordering solution grounded in actual client requirements rather than a purely theoretical model."
            ],
            "Future Researchers": [
                "For future researchers, the study may serve as a reference for similar projects involving small food-service establishments that need lightweight digitalization. It can support further work on browser-based ordering, loyalty implementation, queue management, operational monitoring, and other enhancements suitable for independent local businesses."
            ],
        },
    ),
    (
        "Conceptual Framework",
        [
            "The conceptual framework of this study is anchored on the Input-Process-Output (IPO) model. This model is appropriate because it presents the study as a logical transformation of identified operational needs into a functioning software solution. In the context of this project, the IPO model explains how the cafe's present conditions, user requirements, and technical resources are translated through a structured development process into an ordering-centered Progressive Web App system for Jayd's Cafe.",
            "Input refers to the foundational resources and conditions necessary for the development of the system. These inputs include the current operational setup of Jayd's Cafe, observed bottlenecks in manual counter ordering, the use of physical loyalty cards, recurring order errors related to customization, and uncertainty over ingredient availability. They also include the functional requirements of the primary users, namely customers, baristas, and management. On the technical side, the input component covers the hardware, software, and development environment required to build and test the system, along with the business information that the system must manage such as menu data, pricing, user profiles, order details, payment records, loyalty information, and administrative monitoring data.",
            "Process refers to the development approach used to convert those inputs into a working system. This study adopts an Agile Software Development Life Cycle, particularly an iterative and incremental workflow that allows features to be refined based on operational needs and feedback. The process begins with planning, where the scope of the project and its priorities are defined. It is followed by analysis, which focuses on clarifying user requirements and mapping the existing workflow of Jayd's Cafe. The design stage translates these findings into interface structure, data flow, and system architecture. Development then transforms the design into functional modules such as customer ordering, customization, loyalty support, barista monitoring, and administrative oversight. After development, testing is conducted to verify correctness, usability, security, and reliability. Deployment and review complete the cycle by making the system available for actual use and by gathering insights for refinement.",
            "Output refers to the final system produced by the study. The expected output is a Progressive Web App centered on the order lifecycle of Jayd's Cafe. Its customer-facing side supports menu browsing, order placement, beverage customization, loyalty access, and order-status monitoring. Its support side provides barista and administrative visibility over active orders, menu settings, pricing changes, stock-related awareness, promotional mechanisms, and selected analytics relevant to cafe operations. The output is therefore not merely a digital menu, but a coordinated ordering-centered environment that addresses both customer convenience and internal workflow support.",
            "The relationship among these three components is sequential and interdependent. Without accurate input, the process would not respond to the actual needs of Jayd's Cafe. Without a structured development process, the inputs would remain as isolated observations rather than becoming a usable solution. Without a clearly defined output, the study would fail to demonstrate how software development can directly address the identified business problems. For this reason, the IPO model provides a clear and appropriate logic for the development of the proposed system.",
            "To determine whether the output fulfills the intent of the study, system evaluation will be guided by the ISO 25010 software quality framework and supported by end-user feedback. Through this evaluation, the study will assess whether the resulting system is functionally suitable, reliable, secure, and usable within the real operational setting of Jayd's Cafe.",
        ],
    ),
    (
        "Scope and Limitations",
        {
            "Scope": [
                "This study focuses on the development of an ordering-centered Progressive Web App system for Jayd's Cafe. The system is intended to serve three primary user groups. Customers are given access to a mobile-friendly ordering interface through which they may browse the menu, customize beverages, place orders, monitor order progress, and access loyalty and promotional features. Baristas are provided with order-support visibility that helps them monitor incoming requests, follow queue sequence, and respond more effectively to customization details. Management and authorized administrators are given access to features for menu and price management, promotional monitoring, stock visibility, and selected business analytics.",
                "To meet these users' needs, the scope of the system is organized into four major functional areas.",
            ],
            "Modules and Functions": [
                "User Registration and Profile Module: manages account-related information and user identity within the platform.",
                "Digital Menu, Customization, and Order Management Module: covers browsing, order submission, special remarks, order-status visibility, and payment support through Maya Business or configured counter-based flows.",
                "Digital Loyalty and Promotional Module: replaces the physical stamp-card setup with a digital mechanism for rewards and coupons.",
                "Barista and Administrative Monitoring Module: supports queue visibility, menu and pricing updates, stock-awareness functions, and selected business indicators relevant to operations.",
            ],
            "Limitations": [
                "The system is delivered as a Progressive Web App and is intended for browser-based use; it is not a native mobile application for app-store distribution.",
                "The study does not create a proprietary payment gateway, but instead supports third-party digital payment through Maya Business and related operational payment handling.",
                "The system is not designed to manage payroll, employee scheduling, or broader human-resource functions.",
                "Although it includes stock visibility and availability support, it is not intended to function as a full supply-chain or enterprise inventory-planning platform.",
                "The study is developed specifically for the operational context of Jayd's Cafe and is therefore tailored to the policies, workflow, and service requirements of that business.",
            ],
        },
    ),
    (
        "Definition of Terms",
        [
            {
                "term": "Admin Dashboard",
                "general": "A digital interface that consolidates monitoring and control functions for authorized personnel.",
                "operational": "In this study, it refers to the management-facing area of the Jayd's Cafe system used to review orders, manage menu-related settings, monitor promotions, view stock-related information, and check selected business indicators.",
            },
            {
                "term": "Agile Software Development Life Cycle",
                "general": "An iterative and incremental approach to software development that organizes work into repeatable cycles.",
                "operational": "In this study, it refers to the development approach used by the researchers to plan, build, test, review, and refine the proposed Jayd's Cafe system in stages.",
            },
            {
                "term": "Digital Loyalty System",
                "general": "A technology-based rewards mechanism that tracks and manages customer incentives electronically.",
                "operational": "In this study, it refers to the feature of the proposed system that replaces Jayd's Cafe's physical stamp cards with digital reward and coupon support.",
            },
            {
                "term": "ISO 25010 Software Quality Framework",
                "general": "An international standard used to evaluate software quality through criteria such as functionality, reliability, security, and usability.",
                "operational": "In this study, it refers to the basis for assessing the quality and acceptability of the developed Jayd's Cafe system.",
            },
            {
                "term": "Maya Business",
                "general": "A third-party digital payment service that allows merchants to accept electronic transactions.",
                "operational": "In this study, it refers to the external payment channel integrated into the proposed system to support digital payment handling for Jayd's Cafe orders.",
            },
            {
                "term": "Progressive Web App or PWA",
                "general": "A web application designed to deliver app-like functionality through a standards-compliant browser.",
                "operational": "In this study, it refers to the browser-based ordering platform developed for Jayd's Cafe so customers can access the system without installing a native mobile application.",
            },
            {
                "term": "Role-Based Access Control or RBAC",
                "general": "A security approach that restricts system access according to user roles and permissions.",
                "operational": "In this study, it refers to the access structure that separates customer, barista, and management privileges within the proposed Jayd's Cafe system.",
            },
        ],
    ),
]


def body_word_count() -> int:
    parts = []
    for _, section in BODY_SECTIONS:
        if isinstance(section, list):
            if section and isinstance(section[0], dict):
                for item in section:
                    parts.extend([item["term"], item["general"], item["operational"]])
            else:
                parts.extend(section)
        elif isinstance(section, dict):
            for values in section.values():
                parts.extend(values)
    return len(re.findall(r"\b[\w'-]+\b", "\n\n".join(parts)))


def set_run_font(run, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Open Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Open Sans")
    r_fonts.set(qn("w:hAnsi"), "Open Sans")
    r_fonts.set(qn("w:eastAsia"), "Open Sans")


def format_paragraph(
    paragraph,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_indent: float | None = 0.5,
    left_indent: float | None = None,
    line_spacing: float = 2.0,
    space_before: float = 0,
    space_after: float = 0,
) -> None:
    paragraph.alignment = align
    paragraph_format = paragraph.paragraph_format
    paragraph_format.line_spacing = line_spacing
    paragraph_format.space_before = Pt(space_before)
    paragraph_format.space_after = Pt(space_after)
    paragraph_format.first_line_indent = Inches(first_line_indent) if first_line_indent is not None else None
    paragraph_format.left_indent = Inches(left_indent) if left_indent is not None else None


def apply_font_to_paragraph(paragraph, size: float = 11) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bool(run.bold), italic=bool(run.italic))


def clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def configure_section(section) -> None:
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        configure_section(section)

    normal = next(style for style in doc.styles if (style.name or "").lower() == "normal")
    normal.font.name = "Open Sans"
    normal.font.size = Pt(11)
    normal_r_pr = normal.element.get_or_add_rPr()
    normal_r_fonts = normal_r_pr.get_or_add_rFonts()
    normal_r_fonts.set(qn("w:ascii"), "Open Sans")
    normal_r_fonts.set(qn("w:hAnsi"), "Open Sans")
    normal_r_fonts.set(qn("w:eastAsia"), "Open Sans")


def add_centered_paragraph(doc: Document, text: str, *, size: float = 12, bold: bool = False, after: float = 0) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=None,
        line_spacing=1.0,
        space_after=after,
    )


def add_section_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=12, bold=True)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=6,
        space_after=6,
    )


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    format_paragraph(paragraph)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"{chr(8226)} {text}")
    set_run_font(run, size=11)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=-0.25,
        left_indent=0.5,
        space_after=0,
    )


def add_labeled_paragraph(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(f"{label} ")
    set_run_font(label_run, size=11, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11)
    format_paragraph(paragraph)


def add_caption_label(doc: Document, label: str, number: int) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"{label} {number}")
    set_run_font(run, size=11, bold=True)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=6,
        space_after=0,
    )


def add_caption_title(doc: Document, title: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(title)
    set_run_font(run, size=11, italic=True)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=None,
        line_spacing=1.0,
        space_before=0,
        space_after=6,
    )


def add_title_page(doc: Document) -> None:
    doc.add_paragraph()
    doc.add_paragraph()
    for line in TITLE_LINES:
        add_centered_paragraph(doc, line, size=12, bold=True, after=2)

    doc.add_paragraph()
    add_centered_paragraph(doc, "AN UNDERGRADUATE CAPSTONE PROJECT", size=12, bold=True, after=12)
    add_centered_paragraph(doc, "Presented to the Faculty of the", size=11, after=0)
    add_centered_paragraph(doc, "Institute of Computing and Digital Innovations", size=11, after=0)
    add_centered_paragraph(doc, "Kolehiyo ng Lungsod ng Dasmarinas", size=11, after=0)
    add_centered_paragraph(doc, "City of Dasmarinas, Cavite, Philippines", size=11, after=12)
    add_centered_paragraph(doc, "In Partial Fulfilment", size=11, after=0)
    add_centered_paragraph(doc, "of the Requirements for the Degree", size=11, after=0)
    add_centered_paragraph(doc, "Bachelor of Science in Information Systems", size=11, after=18)
    add_centered_paragraph(doc, "AARON C. MARTINEZ", size=11, after=0)
    add_centered_paragraph(doc, "CHARMAINE G. LLAWDERES", size=11, after=0)
    add_centered_paragraph(doc, "JAMES WILFRED B. ARGA", size=11, after=18)
    add_centered_paragraph(doc, "APRIL 2026", size=11, after=0)


def clear_story_container(container) -> None:
    element = container._element
    for child in list(element):
        element.remove(child)
    container.add_paragraph()


def add_page_number_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, size=11)

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")

    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "

    field_separator = OxmlElement("w:fldChar")
    field_separator.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "1"

    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")

    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_separator)
    run._r.append(placeholder)
    run._r.append(field_end)


def set_section_page_start(section, start: int) -> None:
    section_properties = section._sectPr
    page_number_type = section_properties.find(qn("w:pgNumType"))
    if page_number_type is None:
        page_number_type = OxmlElement("w:pgNumType")
        section_properties.append(page_number_type)
    page_number_type.set(qn("w:start"), str(start))


def configure_section_headers(section, *, show_page_number: bool, hide_first_page: bool = False) -> None:
    section.different_first_page_header_footer = hide_first_page

    for story in [section.header, section.footer, section.first_page_header, section.first_page_footer]:
        story.is_linked_to_previous = False
        clear_story_container(story)

    if show_page_number:
        header_paragraph = section.header.paragraphs[0]
        format_paragraph(
            header_paragraph,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            first_line_indent=None,
            line_spacing=1.0,
        )
        add_page_number_field(header_paragraph)

    blank_header = section.first_page_header.paragraphs[0]
    format_paragraph(
        blank_header,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_line_indent=None,
        line_spacing=1.0,
    )

    for story in [section.footer, section.first_page_footer]:
        footer_paragraph = story.paragraphs[0]
        format_paragraph(
            footer_paragraph,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            first_line_indent=None,
            line_spacing=1.0,
        )


def apply_table_grid(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)

    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")


def shade_cell(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def load_pil_font(size: int):
    from PIL import ImageFont

    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf",
        r"C:\Windows\Fonts\Arialbd.ttf",
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\Arial.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def wrap_text(draw, text: str, font, max_width: int) -> list[str]:
    lines: list[str] = []
    for block in text.split("\n"):
        stripped = block.strip()
        if not stripped:
            lines.append("")
            continue
        words = stripped.split()
        current = words[0]
        for word in words[1:]:
            trial = f"{current} {word}"
            width = draw.textbbox((0, 0), trial, font=font)[2]
            if width <= max_width:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def draw_wrapped_block(draw, bounds, text: str, *, header: bool = False) -> None:
    x1, y1, x2, y2 = bounds
    font = load_pil_font(30 if header else 24)
    lines = wrap_text(draw, text, font, x2 - x1)
    line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + (8 if header else 6)
    current_y = y1
    for line in lines:
        if header:
            line_width = draw.textbbox((0, 0), line, font=font)[2]
            draw.text((x1 + ((x2 - x1 - line_width) / 2), current_y), line, fill="black", font=font)
        else:
            draw.text((x1, current_y), line, fill="black", font=font)
        current_y += line_height
        if current_y > y2:
            break


def generate_research_paradigm_image(output_path: Path) -> None:
    from PIL import Image, ImageDraw

    width = 1600
    height = 1250
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    top = 80
    box_height = 760
    box_width = 360
    gap = 120
    left = 70
    header_height = 70
    stroke = 4

    box_positions = [
        (left, top, left + box_width, top + box_height),
        (left + box_width + gap, top, left + (box_width * 2) + gap, top + box_height),
        (left + (box_width * 2) + (gap * 2), top, left + (box_width * 3) + (gap * 2), top + box_height),
    ]

    titles = ["INPUT", "PROCESS", "OUTPUT"]
    contents = [
        "Client interview findings\nCurrent ordering workflow\nPeak-hour delays and queue buildup\nOrder mix-ups and stock concerns\nFunctional requirements of customers, baristas, and management\nMenu, order, loyalty, and payment data\nHardware and software resources",
        "Requirements gathering\nPlanning and analysis\nInterface, database, and architecture design\nDevelopment and feature integration\nTesting and ISO 25010 evaluation\nRevision and deployment",
        "Ordering-centered Progressive Web App for Jayd's Cafe with:\nCustomer ordering and beverage customization\nPOS and admin monitoring support\nDigital loyalty and promotions\nStock awareness and analytics",
    ]

    for index, (x1, y1, x2, y2) in enumerate(box_positions):
        draw.rectangle((x1, y1, x2, y2), outline="black", width=stroke)
        draw.line((x1, y1 + header_height, x2, y1 + header_height), fill="black", width=stroke)
        draw_wrapped_block(draw, (x1 + 10, y1 + 14, x2 - 10, y1 + header_height - 10), titles[index], header=True)
        draw_wrapped_block(draw, (x1 + 18, y1 + header_height + 16, x2 - 18, y2 - 16), contents[index])

    def draw_arrow_between(start_x: int, mid_y: int) -> None:
        arrow = [
            (start_x, mid_y - 35),
            (start_x + 45, mid_y - 35),
            (start_x + 45, mid_y - 60),
            (start_x + 90, mid_y),
            (start_x + 45, mid_y + 60),
            (start_x + 45, mid_y + 35),
            (start_x, mid_y + 35),
        ]
        draw.polygon(arrow, outline="black", fill="white")

    def draw_left_arrowhead(tip_x: int, center_y: int, *, width: int = 40, height: int = 25) -> None:
        draw.polygon(
            [
                (tip_x, center_y),
                (tip_x + width, center_y - height),
                (tip_x + width, center_y + height),
            ],
            outline="black",
            fill="white",
        )

    def draw_up_arrowhead(center_x: int, tip_y: int, *, width: int = 25, height: int = 40) -> None:
        draw.polygon(
            [
                (center_x, tip_y),
                (center_x - width, tip_y + height),
                (center_x + width, tip_y + height),
            ],
            outline="black",
            fill="white",
        )

    mid_y = top + 360
    draw_arrow_between(box_positions[0][2] + 12, mid_y)
    draw_arrow_between(box_positions[1][2] + 12, mid_y)

    feedback_box = (620, 980, 980, 1050)
    draw.rectangle(feedback_box, outline="black", width=stroke)
    draw_wrapped_block(
        draw,
        (feedback_box[0] + 10, feedback_box[1] + 18, feedback_box[2] - 10, feedback_box[3] - 10),
        "FEEDBACK",
        header=True,
    )

    # Left return path from feedback to input.
    input_center_x = (box_positions[0][0] + box_positions[0][2]) // 2
    input_bottom_y = box_positions[0][3]
    left_feedback_y = (feedback_box[1] + feedback_box[3]) // 2
    left_vertical_y = input_bottom_y + 70
    draw.line((feedback_box[0], left_feedback_y, input_center_x, left_feedback_y), fill="black", width=stroke)
    draw.line((input_center_x, left_feedback_y, input_center_x, left_vertical_y), fill="black", width=stroke)
    draw.line((input_center_x, left_vertical_y, input_center_x, input_bottom_y + 42), fill="black", width=stroke)
    draw_up_arrowhead(input_center_x, input_bottom_y)

    # Right return path from output to feedback.
    output_center_x = (box_positions[2][0] + box_positions[2][2]) // 2
    right_feedback_y = (feedback_box[1] + feedback_box[3]) // 2
    arrow_base_x = feedback_box[2] + 40
    draw.line((output_center_x, box_positions[2][3], output_center_x, right_feedback_y), fill="black", width=stroke)
    draw.line((output_center_x, right_feedback_y, arrow_base_x, right_feedback_y), fill="black", width=stroke)
    draw_left_arrowhead(feedback_box[2], right_feedback_y)

    image.save(output_path)


def add_research_paradigm_figure(doc: Document) -> None:
    generate_research_paradigm_image(FIGURE_PATH)
    add_caption_label(doc, "Figure", 1)
    add_caption_title(doc, "Research Paradigm of the Study")
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(str(FIGURE_PATH), width=Inches(5.9))
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=None,
        line_spacing=1.0,
        space_after=6,
    )


def add_ipo_table(doc: Document) -> None:
    add_caption_label(doc, "Table", 1)
    add_caption_title(doc, "Input-Process-Output (IPO) Conceptual Framework of the Study")

    table = doc.add_table(rows=2, cols=3)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_grid(table)

    headers = ["Input", "Process", "Output"]
    values = [
        "Client interview findings\nCurrent ordering and loyalty workflow\nFunctional requirements of customers, baristas, and management\nHardware, software, and business data resources",
        "Requirements gathering\nPlanning and analysis\nInterface, database, and architecture design\nDevelopment and feature integration\nTesting and ISO 25010 evaluation",
        "Ordering-centered PWA for Jayd's Cafe\nCustomer ordering and beverage customization\nPOS and admin monitoring support\nDigital loyalty, promotions, and stock awareness",
    ]

    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, "D9E2F3")
        for paragraph in cell.paragraphs:
            if paragraph.runs:
                paragraph.runs[0].bold = True
            apply_font_to_paragraph(paragraph, size=11)
            format_paragraph(
                paragraph,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                first_line_indent=None,
                line_spacing=1.0,
            )

    for index, value in enumerate(values):
        cell = table.cell(1, index)
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            apply_font_to_paragraph(paragraph, size=11)
            format_paragraph(
                paragraph,
                align=WD_ALIGN_PARAGRAPH.LEFT,
                first_line_indent=None,
                left_indent=None,
                line_spacing=1.0,
            )


def build_document(output_path: Path, template_path: Path) -> int:
    doc = Document(str(template_path))
    clear_document(doc)
    configure_document(doc)

    title_section = doc.sections[0]
    configure_section_headers(title_section, show_page_number=False)
    add_title_page(doc)

    chapter_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(chapter_section)
    configure_section_headers(chapter_section, show_page_number=True, hide_first_page=True)
    set_section_page_start(chapter_section, 1)

    add_centered_paragraph(doc, "CHAPTER 1", size=12, bold=True, after=0)
    add_centered_paragraph(doc, "INTRODUCTION", size=12, bold=True, after=12)

    for section_name, content in BODY_SECTIONS:
        add_section_heading(doc, section_name)

        if section_name == "Objectives of the Study":
            general = content["General Objective"]
            specific = content["Specific Objectives"]
            add_section_heading(doc, "General Objective")
            for paragraph in general:
                add_body_paragraph(doc, paragraph)
            add_section_heading(doc, "Specific Objectives")
            intro = doc.add_paragraph()
            intro_run = intro.add_run("Specifically, the study aims to:")
            set_run_font(intro_run, size=11)
            format_paragraph(intro)
            for bullet in specific:
                add_bullet(doc, bullet)
            continue

        if section_name == "Conceptual Framework":
            for index, paragraph in enumerate(content):
                add_body_paragraph(doc, paragraph)
                if index == 0:
                    add_research_paradigm_figure(doc)
            continue

        if section_name == "Purpose and Description":
            for label, paragraphs in content.items():
                for paragraph in paragraphs:
                    add_labeled_paragraph(doc, f"{label}:", paragraph)
            continue

        if section_name == "Scope and Limitations":
            for paragraph in content["Scope"]:
                add_body_paragraph(doc, paragraph)
            add_section_heading(doc, "Modules and Functions")
            for bullet in content["Modules and Functions"]:
                add_bullet(doc, bullet)
            add_section_heading(doc, "Limitations")
            for bullet in content["Limitations"]:
                add_bullet(doc, bullet)
            continue

        if section_name == "Definition of Terms":
            for item in content:
                add_section_heading(doc, item["term"])
                add_labeled_paragraph(doc, "General:", item["general"])
                add_labeled_paragraph(doc, "Operational:", item["operational"])
            continue

        for paragraph in content:
            add_body_paragraph(doc, paragraph)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return body_word_count()


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the revised Chapter One .docx file.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    parser.add_argument("--template", type=Path, default=TEMPLATE_PATH)
    parser.add_argument("--dry-run", action="store_true", help="Print the draft word count without writing a file.")
    args = parser.parse_args()

    count = body_word_count()
    if args.dry_run:
        print(f"BODY_WORD_COUNT={count}")
        return

    final_count = build_document(args.output, args.template)
    print(f"BODY_WORD_COUNT={final_count}")
    print(f"WROTE={args.output}")


if __name__ == "__main__":
    main()
