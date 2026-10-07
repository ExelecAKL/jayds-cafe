from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START

from chapter2_docx_generator import (
    SOURCE_PATH,
    add_body_paragraph,
    add_caption_label,
    add_caption_title,
    add_centered_paragraph,
    add_heading,
    add_markdown_table,
    add_reference_paragraph,
    body_word_count,
    configure_section,
    configure_section_headers,
    parse_source,
    reference_count,
)


DEFAULT_BASE_PATH = Path(
    r"C:\Users\Exelec\Downloads\Documents & Data\Documents\MARTINEZ_LLAWDERES_ARGA_C-1DRAFT_REVISED_KLD-STRICT-CAPTIONS.docx"
)
DEFAULT_OUTPUT_PATH = Path(__file__).with_name("MARTINEZ_LLAWDERES_ARGA_CHAPTERS_1_AND_2_REVISED.docx")


def append_chapter_two(doc: Document, source_path: Path) -> None:
    chapter_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(chapter_section)
    configure_section_headers(chapter_section, show_page_number=True, hide_first_page=True)

    add_centered_paragraph(doc, "CHAPTER 2", size=12, bold=True, after=0)
    add_centered_paragraph(doc, "REVIEW OF RELATED LITERATURE AND STUDIES", size=12, bold=True, after=12)

    for token_type, value in parse_source(source_path):
        if token_type == "h2":
            add_heading(doc, value, size=12, after=6)
        elif token_type == "h3":
            add_heading(doc, value, size=11, after=4)
        elif token_type == "h4":
            add_heading(doc, value, size=11, italic=True, after=2)
        elif token_type == "caption":
            caption_type, number, title = value
            add_caption_label(doc, caption_type, number)
            add_caption_title(doc, title)
        elif token_type == "table":
            add_markdown_table(doc, value)
        elif token_type == "ref":
            add_reference_paragraph(doc, value)
        else:
            add_body_paragraph(doc, value)


def build_document(base_path: Path, output_path: Path, source_path: Path) -> tuple[int, int]:
    doc = Document(str(base_path))
    append_chapter_two(doc, source_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return body_word_count(source_path), reference_count(source_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Append the revised Chapter Two to the generated Chapter One manuscript.")
    parser.add_argument("--base", type=Path, default=DEFAULT_BASE_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    parser.add_argument("--source", type=Path, default=SOURCE_PATH)
    parser.add_argument("--dry-run", action="store_true", help="Print Chapter Two body and reference counts without writing a file.")
    args = parser.parse_args()

    if args.dry_run:
        print(f"BODY_WORD_COUNT={body_word_count(args.source)}")
        print(f"REFERENCE_COUNT={reference_count(args.source)}")
        print(f"BASE={args.base}")
        return

    words, refs = build_document(args.base, args.output, args.source)
    print(f"BODY_WORD_COUNT={words}")
    print(f"REFERENCE_COUNT={refs}")
    print(f"WROTE={args.output}")


if __name__ == "__main__":
    main()
